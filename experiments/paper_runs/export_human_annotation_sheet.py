from __future__ import annotations

import argparse
import csv
import json
import os
import re
import subprocess
import time
from pathlib import Path
from typing import Any

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation


DEFAULT_INPUT_JSON = (
    "artifacts/experiment/paper_runs/blind_comment_judge/blind_comment_samples.json"
)
DEFAULT_JUDGED_CSV = (
    "artifacts/experiment/paper_runs/blind_comment_judge/blind_comment_judged.csv"
)
DEFAULT_OUTPUT_DIR = "artifacts/human_annotation/blind_comment_zh"
STANCE_OPTIONS = [
    "Agree（认同/附和）",
    "Disagree（反驳/质疑）",
]
EMOTION_OPTIONS = [
    "skepticism（怀疑）",
    "vigilance（警惕）",
    "anger（愤怒）",
    "panic（恐慌）",
    "fanaticism（狂热）",
    "confusion（困惑）",
    "neutral（中性）",
]
STANCE_LABEL_MAP = {
    "true": STANCE_OPTIONS[0],
    "false": STANCE_OPTIONS[1],
    "agree": STANCE_OPTIONS[0],
    "disagree": STANCE_OPTIONS[1],
}
EMOTION_LABEL_MAP = {
    "skepticism": EMOTION_OPTIONS[0],
    "vigilance": EMOTION_OPTIONS[1],
    "anger": EMOTION_OPTIONS[2],
    "panic": EMOTION_OPTIONS[3],
    "fanaticism": EMOTION_OPTIONS[4],
    "confusion": EMOTION_OPTIONS[5],
    "neutral": EMOTION_OPTIONS[6],
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Translate blind comment samples to Chinese and export a human annotation sheet."
        )
    )
    parser.add_argument("--input-json", default=DEFAULT_INPUT_JSON)
    parser.add_argument("--judged-csv", default=DEFAULT_JUDGED_CSV)
    parser.add_argument("--output-dir", default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--model", default="gpt-4o-mini")
    parser.add_argument("--batch-size", type=int, default=20)
    parser.add_argument("--sleep-seconds", type=float, default=0.5)
    parser.add_argument("--limit", type=int, default=0)
    parser.add_argument("--force", action="store_true")
    return parser.parse_args()


def hydrate_openai_env() -> tuple[str, str]:
    api_key = os.environ.get("OPENAI_API_KEY") or os.environ.get("GPT_DIRECT_API_KEY")
    base_url = (
        os.environ.get("OPENAI_API_BASE_URL")
        or os.environ.get("OPENAI_BASE_URL")
        or "https://api.openai.com/v1"
    )
    if api_key:
        return api_key, base_url.rstrip("/")

    try:
        proc = subprocess.run(
            ["pgrep", "-f", "python .*twitter_simulation_large.py"],
            check=False,
            capture_output=True,
            text=True,
        )
        pids = [line.strip() for line in proc.stdout.splitlines() if line.strip()]
    except OSError:
        pids = []

    for pid in pids:
        environ_path = Path("/proc") / pid / "environ"
        if not environ_path.exists():
            continue
        data = environ_path.read_bytes().split(b"\0")
        values: dict[str, str] = {}
        for entry in data:
            if b"=" not in entry:
                continue
            key, value = entry.split(b"=", 1)
            values[key.decode("utf-8", errors="ignore")] = value.decode(
                "utf-8", errors="ignore"
            )
        api_key = values.get("OPENAI_API_KEY") or values.get("GPT_DIRECT_API_KEY")
        if api_key:
            base_url = (
                values.get("OPENAI_API_BASE_URL")
                or values.get("OPENAI_BASE_URL")
                or base_url
            )
            os.environ.setdefault("OPENAI_API_KEY", api_key)
            os.environ.setdefault("GPT_DIRECT_API_KEY", api_key)
            os.environ.setdefault("OPENAI_API_BASE_URL", base_url)
            return api_key, base_url.rstrip("/")

    raise RuntimeError(
        "No OpenAI-compatible API key found in the environment or reusable simulation processes."
    )


def load_json(path: Path) -> list[dict[str, Any]]:
    return json.loads(path.read_text(encoding="utf-8"))


def load_judged_rows(path: Path) -> dict[str, dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as fh:
        reader = csv.DictReader(fh)
        return {row["sample_id"]: row for row in reader}


def chunks(items: list[dict[str, Any]], size: int) -> list[list[dict[str, Any]]]:
    return [items[i : i + size] for i in range(0, len(items), size)]


def extract_json_array(text: str) -> list[dict[str, Any]]:
    text = text.strip()
    fenced = re.search(r"```(?:json)?\s*(\[.*\])\s*```", text, flags=re.S)
    if fenced:
        text = fenced.group(1)
    if not text.startswith("["):
        start = text.find("[")
        end = text.rfind("]")
        if start != -1 and end != -1 and end > start:
            text = text[start : end + 1]
    data = json.loads(text)
    if not isinstance(data, list):
        raise ValueError("Expected a JSON array.")
    return data


def translate_batch(
    batch: list[dict[str, Any]],
    *,
    api_key: str,
    base_url: str,
    model: str,
    timeout_seconds: int = 120,
) -> list[dict[str, str]]:
    payload_rows = [
        {
            "sample_id": item["sample_id"],
            "post_content": item["post_content"],
            "comment_content": item["comment_content"],
        }
        for item in batch
    ]
    system_prompt = (
        "You are a precise bilingual translator. Translate English social-media misinformation posts "
        "and comments into Simplified Chinese. Preserve meaning, stance, intensity, and misinformation "
        "content faithfully. Do not sanitize, correct, or add warnings. Return only a JSON array. "
        "Each object must contain sample_id, post_zh, comment_zh."
    )
    user_prompt = json.dumps(payload_rows, ensure_ascii=False)
    response = requests.post(
        f"{base_url}/chat/completions",
        headers={
            "Authorization": f"Bearer {api_key}",
            "Content-Type": "application/json",
        },
        json={
            "model": model,
            "temperature": 0,
            "messages": [
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": user_prompt},
            ],
        },
        timeout=timeout_seconds,
    )
    response.raise_for_status()
    data = response.json()
    content = data["choices"][0]["message"]["content"]
    rows = extract_json_array(content)
    result: list[dict[str, str]] = []
    for row in rows:
        result.append(
            {
                "sample_id": str(row["sample_id"]),
                "post_zh": str(row["post_zh"]).strip(),
                "comment_zh": str(row["comment_zh"]).strip(),
            }
        )
    return result


def load_cache(path: Path) -> dict[str, dict[str, str]]:
    if not path.exists():
        return {}
    data = json.loads(path.read_text(encoding="utf-8"))
    return {str(row["sample_id"]): row for row in data}


def save_cache(path: Path, cache: dict[str, dict[str, str]]) -> None:
    rows = [cache[key] for key in sorted(cache, key=lambda x: int(x))]
    path.write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8")


def ensure_translations(
    samples: list[dict[str, Any]],
    *,
    output_dir: Path,
    model: str,
    batch_size: int,
    sleep_seconds: float,
    force: bool,
) -> dict[str, dict[str, str]]:
    cache_path = output_dir / "translations_cache.json"
    cache = {} if force else load_cache(cache_path)
    missing = [item for item in samples if str(item["sample_id"]) not in cache]
    if not missing:
        return cache

    api_key, base_url = hydrate_openai_env()
    grouped = chunks(missing, batch_size)
    for idx, batch in enumerate(grouped, start=1):
        translated = translate_batch(
            batch,
            api_key=api_key,
            base_url=base_url,
            model=model,
        )
        for row in translated:
            cache[str(row["sample_id"])] = row
        save_cache(cache_path, cache)
        print(f"translated batch {idx}/{len(grouped)}")
        time.sleep(sleep_seconds)
    return cache


def add_instructions_sheet(wb: Workbook) -> None:
    ws = wb.active
    ws.title = "标注说明"
    lines = [
        "评论人工标注说明",
        "",
        "1. 请在“评论标注”工作表中填写，不要修改 sample_id。",
        "2. “人工标注立场（下拉选择）”列请直接从下拉菜单里选 1 个，不要手输。",
        "   选项：Agree（认同/附和） 或 Disagree（反驳/质疑）。",
        "3. “人工标注情绪（下拉选择）”列请直接从下拉菜单里选 1 个，不要手输。",
        "   选项：skepticism（怀疑）、vigilance（警惕）、anger（愤怒）、panic（恐慌）、fanaticism（狂热）、confusion（困惑）、neutral（中性）。",
        "4. 主表只保留中文内容；如需对照模型判断，请查看“模型结果参考”工作表。",
        "5. 如遇表达模糊或难判样本，可在“备注”列补充说明。",
        "6. 如果没有看到下拉箭头，请先单击单元格。",
        "7. 该表故意隐藏实验条件信息，便于盲标。",
    ]
    for row_idx, line in enumerate(lines, start=1):
        ws.cell(row=row_idx, column=1, value=line)
    ws.column_dimensions["A"].width = 100
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws["A1"].font = Font(bold=True, size=14)


def build_dropdown_formula(options: list[str]) -> str:
    escaped = [option.replace('"', '""') for option in options]
    return '"' + ",".join(escaped) + '"'


def add_dropdown_validations(
    ws,
    last_row: int,
    *,
    stance_col: str,
    emotion_col: str,
) -> None:
    if last_row < 2:
        return

    stance_validation = DataValidation(
        type="list",
        formula1=build_dropdown_formula(STANCE_OPTIONS),
        allow_blank=True,
    )
    stance_validation.promptTitle = "选择立场"
    stance_validation.prompt = "请从下拉菜单中选择 1 个立场标签。"
    stance_validation.errorTitle = "无效选项"
    stance_validation.error = "请从下拉菜单中选择立场，不要手动输入其他内容。"

    emotion_validation = DataValidation(
        type="list",
        formula1=build_dropdown_formula(EMOTION_OPTIONS),
        allow_blank=True,
    )
    emotion_validation.promptTitle = "选择情绪"
    emotion_validation.prompt = "请从下拉菜单中选择 1 个情绪标签。"
    emotion_validation.errorTitle = "无效选项"
    emotion_validation.error = "请从下拉菜单中选择情绪，不要手动输入其他内容。"

    ws.add_data_validation(stance_validation)
    ws.add_data_validation(emotion_validation)
    stance_validation.add(f"{stance_col}2:{stance_col}{last_row}")
    emotion_validation.add(f"{emotion_col}2:{emotion_col}{last_row}")


def normalize_stance_label(value: str) -> str:
    key = str(value).strip().lower()
    return STANCE_LABEL_MAP.get(key, str(value).strip())


def normalize_emotion_label(value: str) -> str:
    key = str(value).strip().lower()
    return EMOTION_LABEL_MAP.get(key, str(value).strip())


def write_annotation_workbook(
    annotation_rows: list[dict[str, Any]],
    llm_rows: list[dict[str, Any]],
    output_path: Path,
) -> None:
    wb = Workbook()
    add_instructions_sheet(wb)
    ws = wb.create_sheet("评论标注")
    headers = [
        "sample_id",
        "帖子中文",
        "评论中文",
        "人工标注立场（下拉选择）",
        "人工标注情绪（下拉选择）",
        "备注",
    ]
    ws.append(headers)
    for row in annotation_rows:
        ws.append(
            [
                row["sample_id"],
                row["post_zh"],
                row["comment_zh"],
                "",
                "",
                "",
            ]
        )
    widths = {
        "A": 12,
        "B": 55,
        "C": 55,
        "D": 18,
        "E": 18,
        "F": 28,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    add_dropdown_validations(ws, ws.max_row, stance_col="D", emotion_col="E")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    ref_ws = wb.create_sheet("模型结果参考")
    ref_headers = [
        "sample_id",
        "帖子中文",
        "评论中文",
        "大模型标注立场",
        "大模型标注情绪",
        "大模型置信度",
    ]
    ref_ws.append(ref_headers)
    for row in llm_rows:
        ref_ws.append(
            [
                row["sample_id"],
                row["post_zh"],
                row["comment_zh"],
                row["llm_stance"],
                row["llm_emotion"],
                row["llm_confidence"],
            ]
        )
    ref_widths = {
        "A": 12,
        "B": 55,
        "C": 55,
        "D": 20,
        "E": 24,
        "F": 14,
    }
    for col, width in ref_widths.items():
        ref_ws.column_dimensions[col].width = width
    for cell in ref_ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in ref_ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ref_ws.freeze_panes = "A2"
    ref_ws.auto_filter.ref = ref_ws.dimensions
    wb.save(output_path)


def write_csv(path: Path, headers: list[str], rows: list[dict[str, Any]]) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.DictWriter(fh, fieldnames=headers)
        writer.writeheader()
        writer.writerows(rows)


def build_rows(
    samples: list[dict[str, Any]],
    judged_map: dict[str, dict[str, str]],
    translations: dict[str, dict[str, str]],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    annotation_rows: list[dict[str, Any]] = []
    llm_rows: list[dict[str, Any]] = []
    key_rows: list[dict[str, Any]] = []
    for item in sorted(samples, key=lambda x: int(x["sample_id"])):
        sample_id = str(item["sample_id"])
        trans = translations[sample_id]
        judged = judged_map.get(sample_id, {})
        annotation_row = {
            "sample_id": sample_id,
            "post_zh": trans["post_zh"],
            "comment_zh": trans["comment_zh"],
            "human_stance": "",
            "human_emotion": "",
            "notes": "",
        }
        llm_row = {
            "sample_id": sample_id,
            "post_zh": trans["post_zh"],
            "comment_zh": trans["comment_zh"],
            "llm_stance": normalize_stance_label(judged.get("judge_agree", "")),
            "llm_emotion": normalize_emotion_label(judged.get("judge_emotion", "")),
            "llm_confidence": judged.get("judge_confidence", ""),
        }
        key_row = {
            **annotation_row,
            "condition": item["condition"],
            "seed": item["seed"],
            "comment_id": item["comment_id"],
            "post_id": item["post_id"],
            "comment_user_id": item["comment_user_id"],
            "post_user_id": item["post_user_id"],
            "post_content": item["post_content"],
            "comment_content": item["comment_content"],
            "self_agree": item["self_agree"],
            "self_emotion": item["self_emotion"],
            "judge_agree": judged.get("judge_agree", ""),
            "judge_emotion": judged.get("judge_emotion", ""),
            "judge_confidence": judged.get("judge_confidence", ""),
        }
        annotation_rows.append(annotation_row)
        llm_rows.append(llm_row)
        key_rows.append(key_row)
    return annotation_rows, llm_rows, key_rows


def main() -> None:
    args = parse_args()
    input_json = Path(args.input_json)
    judged_csv = Path(args.judged_csv)
    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    samples = load_json(input_json)
    if args.limit > 0:
        samples = samples[: args.limit]
    judged_map = load_judged_rows(judged_csv)
    translations = ensure_translations(
        samples,
        output_dir=output_dir,
        model=args.model,
        batch_size=args.batch_size,
        sleep_seconds=args.sleep_seconds,
        force=args.force,
    )
    annotation_rows, llm_rows, key_rows = build_rows(samples, judged_map, translations)

    workbook_path = output_dir / "blind_comment_annotation_sheet_zh.xlsx"
    blind_csv_path = output_dir / "blind_comment_annotation_sheet_zh.csv"
    llm_csv_path = output_dir / "blind_comment_llm_results_zh.csv"
    key_csv_path = output_dir / "blind_comment_annotation_key_internal.csv"
    manifest_path = output_dir / "annotation_manifest.json"

    write_annotation_workbook(annotation_rows, llm_rows, workbook_path)
    write_csv(
        blind_csv_path,
        [
            "sample_id",
            "post_zh",
            "comment_zh",
            "human_stance",
            "human_emotion",
            "notes",
        ],
        annotation_rows,
    )
    write_csv(
        llm_csv_path,
        [
            "sample_id",
            "post_zh",
            "comment_zh",
            "llm_stance",
            "llm_emotion",
            "llm_confidence",
        ],
        llm_rows,
    )
    write_csv(
        key_csv_path,
        [
            "sample_id",
            "condition",
            "seed",
            "comment_id",
            "post_id",
            "comment_user_id",
            "post_user_id",
            "post_content",
            "post_zh",
            "comment_content",
            "comment_zh",
            "self_agree",
            "self_emotion",
            "judge_agree",
            "judge_emotion",
            "judge_confidence",
            "human_stance",
            "human_emotion",
            "notes",
        ],
        key_rows,
    )
    manifest_path.write_text(
        json.dumps(
            {
                "sample_count": len(samples),
                "workbook_path": str(workbook_path),
                "blind_csv_path": str(blind_csv_path),
                "llm_csv_path": str(llm_csv_path),
                "key_csv_path": str(key_csv_path),
                "model": args.model,
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    print(json.dumps(json.loads(manifest_path.read_text(encoding="utf-8")), ensure_ascii=False))


if __name__ == "__main__":
    main()
